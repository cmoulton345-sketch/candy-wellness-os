import sys
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def build_master_ap_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Theme colors (Deep Magenta / Crimson)
    CRIMSON = RGBColor(136, 14, 79)       # #880e4f
    ROSE_LIGHT = RGBColor(252, 228, 236)  # background light
    DARK_NAVY = RGBColor(33, 33, 33)
    WHITE = RGBColor(255, 255, 255)
    CARD_BG = RGBColor(250, 245, 248)
    CARD_BORDER = RGBColor(220, 180, 200)
    ACCENT_MAGENTA = RGBColor(194, 24, 91)

    def add_header(slide, title_text, category_text="HCSS1087A • ANATOMY & PHYSIOLOGY MASTER DECK"):
        header = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(1.2))
        header.fill.solid()
        header.fill.fore_color.rgb = CRIMSON
        header.line.color.rgb = CRIMSON

        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.15), Inches(11.7), Inches(0.35))
        tf_cat = cat_box.text_frame
        tf_cat.word_wrap = True
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category_text.upper()
        p_cat.font.size = Pt(11)
        p_cat.font.bold = True
        p_cat.font.color.rgb = RGBColor(248, 187, 208)

        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.45), Inches(11.7), Inches(0.65))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        p_title = tf_title.paragraphs[0]
        p_title.text = title_text
        p_title.font.size = Pt(24)
        p_title.font.bold = True
        p_title.font.color.rgb = WHITE

    # --- SLIDE 1: Title Slide ---
    slide1 = prs.slides.add_slide(blank_layout)
    bg1 = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = CRIMSON

    tbox1 = slide1.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(11.333), Inches(4.0))
    tf1 = tbox1.text_frame
    tf1.word_wrap = True
    
    p1 = tf1.paragraphs[0]
    p1.text = "HCSS1087A — Anatomy & Physiology"
    p1.font.size = Pt(20)
    p1.font.bold = True
    p1.font.color.rgb = RGBColor(248, 187, 208)

    p2 = tf1.add_paragraph()
    p2.text = "Master Exam Review: Modules 1, 3, 4 & Cell"
    p2.font.size = Pt(38)
    p2.font.bold = True
    p2.font.color.rgb = WHITE
    p2.space_before = Pt(10)

    p3 = tf1.add_paragraph()
    p3.text = "Body Intro & Homeostasis • Cell Organelles • Integumentary System • Skeletal System"
    p3.font.size = Pt(20)
    p3.font.color.rgb = RGBColor(252, 228, 236)
    p3.space_before = Pt(15)

    p4 = tf1.add_paragraph()
    p4.text = "Instructor: Jessica Freeze Snyder | Practical Nursing Master Study Suite for Rayne Moulton"
    p4.font.size = Pt(16)
    p4.font.italic = True
    p4.font.color.rgb = RGBColor(244, 143, 177)
    p4.space_before = Pt(35)

    # --- SLIDE 2: Module 1 Intro & Homeostasis ---
    slide2 = prs.slides.add_slide(blank_layout)
    add_header(slide2, "Module 1: Body Organization, Homeostasis & Cavities")

    m1_cards = [
        ("6 Levels of Organization", [
            "Chemical: Atoms & molecules",
            "Cellular: Basic unit of life",
            "Tissue: Groups of similar cells",
            "Organ: Tissues working together",
            "Organ System: Group of organs",
            "Organism: Complete living human"
        ]),
        ("Homeostasis & Feedback", [
            "Homeostasis: Relative internal body balance",
            "Negative Feedback: Opposes/negates change to restore balance (e.g. temp, glucose)",
            "Positive Feedback: Amplifies change (e.g. child labor, blood clotting)",
            "Loop components: Sensor → Control Center → Effector"
        ]),
        ("Body Cavities & Regions", [
            "Ventral (Front): Thoracic (Mediastinum & Pleural) + Abdominopelvic",
            "Dorsal (Back): Cranial (brain) & Spinal (spinal cord)",
            "4 Quadrants: RUQ, LUQ, RLQ (appendix), LLQ",
            "9 Regions: Epigastric, Umbilical, Hypogastric, Hypochondriac, Flank, Iliac"
        ])
    ]

    for i, (title, bullets) in enumerate(m1_cards):
        left = Inches(0.8 + i * 3.9)
        top = Inches(1.6)
        width = Inches(3.6)
        height = Inches(5.3)

        card = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = CARD_BORDER

        tb = slide2.shapes.add_textbox(left + Inches(0.2), top + Inches(0.2), width - Inches(0.4), height - Inches(0.4))
        tf = tb.text_frame
        tf.word_wrap = True
        
        pt = tf.paragraphs[0]
        pt.text = title
        pt.font.size = Pt(17)
        pt.font.bold = True
        pt.font.color.rgb = CRIMSON

        for b in bullets:
            pb = tf.add_paragraph()
            pb.text = "• " + b
            pb.font.size = Pt(13)
            pb.font.color.rgb = DARK_NAVY
            pb.space_before = Pt(8)

    # --- SLIDE 3: Cell Biology ---
    slide3 = prs.slides.add_slide(blank_layout)
    add_header(slide3, "Cell Biology: Human Cell Structure & Organelles")

    cell_organelles = [
        ("Plasma Membrane", "Semipermeable lipid bilayer enclosing the cytoplasm. Keeps cell intact and regulates entry/exit."),
        ("Nucleus", "Cell's 'control center' managing genetic material (DNA in chromosomes). Directs cell growth & division."),
        ("Ribosomes", "'Protein factories' of the cell. Synthesize proteins; float free in cytoplasm or attach to Endoplasmic Reticulum (ER)."),
        ("Mitochondria", "'Powerhouse' of the cell. Extracts energy from nutrients & converts it into ATP chemical energy.")
    ]

    for idx, (c_name, c_desc) in enumerate(cell_organelles):
        col = idx % 2
        row = idx // 2
        left = Inches(0.8 + col * 5.9)
        top = Inches(1.6 + row * 2.7)
        width = Inches(5.6)
        height = Inches(2.4)

        card = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = CARD_BORDER

        tb = slide3.shapes.add_textbox(left + Inches(0.2), top + Inches(0.2), width - Inches(0.4), height - Inches(0.4))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = c_name
        p.font.size = Pt(18)
        p.font.bold = True
        p.font.color.rgb = ACCENT_MAGENTA

        p2 = tf.add_paragraph()
        p2.text = c_desc
        p2.font.size = Pt(14)
        p2.font.color.rgb = DARK_NAVY
        p2.space_before = Pt(8)

    # --- SLIDE 4: Module 3 Integumentary System ---
    slide4 = prs.slides.add_slide(blank_layout)
    add_header(slide4, "Module 3: Integumentary System (Skin, Glands, Receptors)")

    m3_cards = [
        ("Skin Layers", [
            "Epidermis: Nonvascular; 5 layers (Basale to Corneum). 95% Keratinocytes, 5% Melanocytes.",
            "Dermis: 20x thicker; Papillary (fingerprints) & Reticular (collagen, elastin, blood vessels).",
            "Hypodermis: Adipose fat tissue for insulation & cushioning."
        ]),
        ("Accessory Glands", [
            "Eccrine Sweat Glands: Abundant; thin watery sweat for cooling & waste excretion.",
            "Apocrine Sweat Glands: Axilla & groin; thicker fluid producing odor via skin bacteria.",
            "Sebaceous Glands: Sebum oil prevents skin cracking; blocked ducts form pimples/blackheads."
        ]),
        ("Sensory Receptors", [
            "Meissner's: Light touch & texture",
            "Pacinian: Deep pressure & vibration",
            "Krause: Cold temperature",
            "Ruffini: Skin stretch",
            "Free Nerve Endings: Pain & itch"
        ])
    ]

    for i, (title, bullets) in enumerate(m3_cards):
        left = Inches(0.8 + i * 3.9)
        top = Inches(1.6)
        width = Inches(3.6)
        height = Inches(5.3)

        card = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = CARD_BORDER

        tb = slide4.shapes.add_textbox(left + Inches(0.2), top + Inches(0.2), width - Inches(0.4), height - Inches(0.4))
        tf = tb.text_frame
        tf.word_wrap = True
        
        pt = tf.paragraphs[0]
        pt.text = title
        pt.font.size = Pt(17)
        pt.font.bold = True
        pt.font.color.rgb = CRIMSON

        for b in bullets:
            pb = tf.add_paragraph()
            pb.text = "• " + b
            pb.font.size = Pt(13)
            pb.font.color.rgb = DARK_NAVY
            pb.space_before = Pt(8)

    # --- SLIDE 5: Module 4 Skeletal System & Joints ---
    slide5 = prs.slides.add_slide(blank_layout)
    add_header(slide5, "Module 4: Skeletal System, Cartilage & Joints")

    skel_cards = [
        ("Functions & Bone Types", [
            "5 Functions: Support, Protection, Movement, Calcium Storage, Hematopoiesis (red marrow).",
            "5 Bone Types: Long (femur), Short (carpals), Flat (scapula/sternum), Irregular (vertebrae), Sesamoid (patella)."
        ]),
        ("Skeleton Divisions", [
            "Axial (80 bones): Skull (8 cranial, 14 facial), Hyoid, Spine (26 vertebrae), Thorax (Sternum & 12 pairs of ribs).",
            "Appendicular (126 bones): Upper & lower extremities, Pectoral & Pelvic girdles."
        ]),
        ("Cartilage & Joint Classes", [
            "Cartilage: Hyaline (articular cushion), Elastic (ear), Fibrocartilage (discs/pubis).",
            "Joints: Synarthroses (immovable), Amphiarthroses (slight movement), Diarthroses (freely movable synovial)."
        ])
    ]

    for i, (title, bullets) in enumerate(skel_cards):
        left = Inches(0.8 + i * 3.9)
        top = Inches(1.6)
        width = Inches(3.6)
        height = Inches(5.3)

        card = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = CARD_BORDER

        tb = slide5.shapes.add_textbox(left + Inches(0.2), top + Inches(0.2), width - Inches(0.4), height - Inches(0.4))
        tf = tb.text_frame
        tf.word_wrap = True
        
        pt = tf.paragraphs[0]
        pt.text = title
        pt.font.size = Pt(17)
        pt.font.bold = True
        pt.font.color.rgb = CRIMSON

        for b in bullets:
            pb = tf.add_paragraph()
            pb.text = "• " + b
            pb.font.size = Pt(13)
            pb.font.color.rgb = DARK_NAVY
            pb.space_before = Pt(8)

    # --- SLIDE 6: Summary & Exam Rules ---
    slide6 = prs.slides.add_slide(blank_layout)
    add_header(slide6, "Top A&P Exam Takeaways & Rules for Rayne")

    box_s = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.6), Inches(11.733), Inches(5.3))
    box_s.fill.solid()
    box_s.fill.fore_color.rgb = ROSE_LIGHT
    box_s.line.color.rgb = CRIMSON

    tb_s = slide6.shapes.add_textbox(Inches(1.1), Inches(1.8), Inches(11.1), Inches(4.9))
    tf_s = tb_s.text_frame
    tf_s.word_wrap = True

    p = tf_s.paragraphs[0]
    p.text = "Master A&P Exam Rules:"
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = CRIMSON

    tips = [
        "Rule 1: Negative Feedback = restores balance; Positive Feedback = amplifies change.",
        "Rule 2: Epidermis is NONVASCULAR; receives nutrients from dermal blood vessel diffusion.",
        "Rule 3: Hematopoiesis (blood cell production) occurs in RED bone marrow, NOT yellow marrow.",
        "Rule 4: Hyoid bone is the ONLY bone in the human body that does not articulate with another bone.",
        "Rule 5: True Ribs (1-7) attach directly to sternum; False Ribs (8-10) attach to 7th rib cartilage; Floating Ribs (11-12) unattached anteriorly.",
        "Rule 6: Synarthrotic = Immovable; Amphiarthrotic = Slightly Movable; Diarthrotic = Freely Movable Synovial Joints."
    ]

    for tip in tips:
        pt = tf_s.add_paragraph()
        pt.text = "✔  " + tip
        pt.font.size = Pt(14)
        pt.font.bold = True
        pt.font.color.rgb = DARK_NAVY
        pt.space_before = Pt(10)

    output_path = "/Users/valuedcustomer/Downloads/flowstate_ai_os_candy/Rayne_NBCC/Rayne_AP_Comprehensive_Modules_1_3_4.pptx"
    prs.save(output_path)
    print("Saved Master A&P PowerPoint to:", output_path)

if __name__ == '__main__':
    build_master_ap_deck()
