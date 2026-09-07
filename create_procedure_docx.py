import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=144, bottom=144, left=216, right=216): # values in twips (1/20 of a pt)
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def set_cell_borders(cell, left_color="C00000", left_size="36"): # left_size 36 = 4.5 pt
    tcPr = cell._element.get_or_add_tcPr()
    borders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="none"/>
            <w:left w:val="single" w:sz="{left_size}" w:space="0" w:color="{left_color}"/>
            <w:bottom w:val="none"/>
            <w:right w:val="none"/>
        </w:tcBorders>
    ''')
    tcPr.append(borders)

def add_callout_box(doc, title, text, bg_color="FFF3CD", border_color="856404"):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    cell = table.rows[0].cells[0]
    cell.width = Inches(6.5)
    
    set_cell_background(cell, bg_color)
    set_cell_borders(cell, left_color=border_color, left_size="36")
    set_cell_margins(cell, top=160, bottom=160, left=240, right=240)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    
    run_title = p.add_run(f"{title}\n")
    run_title.font.bold = True
    run_title.font.size = Pt(11)
    run_title.font.name = 'Arial'
    
    # Check if border color is dark red/danger to color title
    if border_color == "C00000":
        run_title.font.color.rgb = RGBColor(192, 0, 0)
    elif border_color == "856404":
        run_title.font.color.rgb = RGBColor(133, 100, 4)
    elif border_color == "0C5460":
        run_title.font.color.rgb = RGBColor(12, 84, 96)
        
    run_text = p.add_run(text)
    run_text.font.size = Pt(10.5)
    run_text.font.name = 'Arial'
    run_text.font.color.rgb = RGBColor(51, 51, 51)
    
    # Add an empty spacing paragraph after callout
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_after = Pt(6)

def add_heading(doc, text, level):
    h = doc.add_heading(text, level=level)
    h.paragraph_format.keep_with_next = True
    h.paragraph_format.space_before = Pt(14 if level==1 else (10 if level==2 else 8))
    h.paragraph_format.space_after = Pt(6 if level==1 else 4)
    for run in h.runs:
        run.font.name = 'Arial'
        if level == 1:
            run.font.size = Pt(15)
            run.font.bold = True
            run.font.color.rgb = RGBColor(26, 54, 93) # Deep navy
        elif level == 2:
            run.font.size = Pt(13)
            run.font.bold = True
            run.font.color.rgb = RGBColor(43, 108, 176) # Slate blue
        elif level == 3:
            run.font.size = Pt(11.5)
            run.font.bold = True
            run.font.color.rgb = RGBColor(45, 55, 72)
    return h

def add_bullet(doc, bold_prefix, text):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    
    if bold_prefix:
        r_bold = p.add_run(bold_prefix)
        r_bold.font.bold = True
        r_bold.font.name = 'Arial'
        r_bold.font.size = Pt(11)
        r_bold.font.color.rgb = RGBColor(30, 41, 59)
        
    r_text = p.add_run(text)
    r_text.font.name = 'Arial'
    r_text.font.size = Pt(11)
    r_text.font.color.rgb = RGBColor(51, 51, 51)
    return p

def add_body(doc, text, bold_prefix=""):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    
    if bold_prefix:
        r_bold = p.add_run(bold_prefix)
        r_bold.font.bold = True
        r_bold.font.name = 'Arial'
        r_bold.font.size = Pt(11)
        r_bold.font.color.rgb = RGBColor(30, 41, 59)
        
    r_text = p.add_run(text)
    r_text.font.name = 'Arial'
    r_text.font.size = Pt(11)
    r_text.font.color.rgb = RGBColor(51, 51, 51)
    return p

def create_document():
    doc = docx.Document()
    
    # Page setup - 1 inch margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
        # Add footer
        footer = section.footer
        f_p = footer.paragraphs[0]
        f_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        f_run = f_p.add_run("Interim Emergency Evacuation Headcount Procedure  |  Company-Provided Housing")
        f_run.font.name = 'Arial'
        f_run.font.size = Pt(8.5)
        f_run.font.color.rgb = RGBColor(128, 128, 128)

    # Document Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(12)
    p_title.paragraph_format.space_after = Pt(2)
    r_title = p_title.add_run("INTERIM EMERGENCY EVACUATION HEADCOUNT PROCEDURE")
    r_title.font.name = 'Arial'
    r_title.font.size = Pt(18)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(26, 54, 93)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(16)
    r_sub = p_sub.add_run("Company-Provided Housing (Swipe-In / Swipe-Out Facilities)")
    r_sub.font.name = 'Arial'
    r_sub.font.size = Pt(13)
    r_sub.font.bold = True
    r_sub.font.color.rgb = RGBColor(100, 116, 139)

    # Interim Warning Box
    add_callout_box(
        doc,
        "INTERIM PROCEDURE — EFFECTIVE IMMEDIATELY",
        "This procedure is an Interim Safety Directive issued to establish positive accountability for all personnel residing in company-provided housing during an emergency evacuation. This procedure remains in full effect until further notice or until superseded by a permanent automated/site-wide headcount system.",
        bg_color="FFF3CD",
        border_color="856404"
    )

    # 1. Purpose
    add_heading(doc, "1. PURPOSE", level=1)
    add_body(doc, "The purpose of this procedure is to establish an immediate, reliable, and standardized process for accounting for all residents and workers evacuated from company-provided accommodation during a fire, gas release, structural emergency, or site evacuation.")

    # 2. Scope
    add_heading(doc, "2. SCOPE", level=1)
    add_body(doc, "This procedure applies across all residential and administrative operations, specifically:")
    add_bullet(doc, "All Residents & Visitors: ", "All employees, contractors, and visitors residing in or utilizing company-provided housing where electronic swipe-in / swipe-out access is utilized.")
    add_bullet(doc, "Supervision: ", "All group supervisors, foremen, and crew leads overseeing personnel residing in housing.")
    add_bullet(doc, "Muster Coordinators: ", "Designated emergency accountability leads assigned to each work group or residential pod.")
    add_bullet(doc, "Emergency Response: ", "On-site Security and Emergency Response Teams (Fire Department / Incident Command).")

    # 3. Regulatory Basis
    add_heading(doc, "3. REGULATORY & STANDARDS BASIS", level=1)
    add_bullet(doc, "WorkSafeBC OHS Regulation Part 4: ", "General Conditions — Emergency Preparedness and Response (Sections 4.13 to 4.18: Emergency evacuation, verification of worker accountability, and training).")
    add_bullet(doc, "CSA Z1000: ", "Occupational Health and Safety Management (Emergency preparedness and response framework).")
    add_bullet(doc, "CSA Z246.1: ", "Security Management for Petroleum and Natural Gas Industry Systems.")

    # 4. Roles & Responsibilities
    add_heading(doc, "4. ROLES & RESPONSIBILITIES", level=1)

    add_heading(doc, "4.1 On-Site Security", level=2)
    add_callout_box(
        doc,
        "SECURITY ROLE CLARIFICATION",
        "Security personnel are stationed at housing evacuation routes and muster areas exclusively to provide physical direction, crowd control, and traffic management.\n\nSecurity does NOT conduct, compile, or verify the personnel headcount. That responsibility rests entirely with the individual work groups through their assigned Muster Coordinators.",
        bg_color="EBF8FF",
        border_color="2B6CB0"
    )

    add_heading(doc, "4.2 Supervisors / Crew Leads", level=2)
    add_bullet(doc, "Roster & Group Chat Maintenance: ", "Ensure a dedicated mobile text messaging group chat (e.g., SMS, WhatsApp, or approved company messaging platform) is established for every group residing in housing.")
    add_bullet(doc, "Continuous Accuracy: ", "Verify and update names, room assignments, and active phone numbers in the group chat immediately whenever personnel rotate on/off shift or change accommodations.")
    add_bullet(doc, "Muster Coordinator Assignment: ", "Ensure a Muster Coordinator (and one backup coordinator) is formally designated at the beginning of each shift cycle or weekly rotation.")

    add_heading(doc, "4.3 Muster Coordinator", level=2)
    add_bullet(doc, "Term of Responsibility: ", "Holds the accountability role for one (1) week at a time, or until relieved, rotated off-shift, or unable to perform duties.")
    add_bullet(doc, "Initiate Check-In: ", "Sends the check-in prompt to the group text chat immediately upon reaching the designated safe muster point during an alarm or evacuation.")
    add_bullet(doc, "Compile & Report: ", "Compiles the headcount and reports missing or unaccounted individuals directly to the Fire Department / Incident Commander / Security Lead within ten (10) minutes of alarm activation.")
    add_bullet(doc, "Muster Discipline: ", "Enforces group cohesion and prevents anyone from leaving the muster area or re-entering the facility.")

    add_heading(doc, "4.4 All Housing Residents / Group Members", level=2)
    add_bullet(doc, "Swipe Discipline: ", "Must personally swipe in and swipe out at all times when entering or exiting housing. Tailgating is strictly prohibited.")
    add_bullet(doc, "Evacuate Immediately: ", "Evacuate immediately upon alarm. Never delay evacuation to send text messages or gather belongings.")
    add_bullet(doc, "Mandatory Check-In: ", "Immediately check in via the group text chat once safely arrived at the designated physical muster point, or report verbally to the Muster Coordinator if phone/battery is unavailable.")
    add_bullet(doc, "Remain Mustered: ", "Stay at the designated muster location until officially released or redirected by Emergency Services.")

    # 5. Procedure
    add_heading(doc, "5. PROCEDURE", level=1)

    add_heading(doc, "Step 1: Pre-Incident Setup & Weekly Rotation", level=2)
    add_bullet(doc, "1. Group Chat Creation: ", "Each operational or residential group creates a dedicated text chat group containing every active resident in that group.")
    add_bullet(doc, "2. Weekly Coordinator Designation: ", "At the start of each work week or shift cycle, the Supervisor designates the Muster Coordinator and a Backup Coordinator.")
    add_bullet(doc, "3. Roster Verification: ", "The Supervisor confirms all active group members have confirmed receipt of a test message in the group chat upon arrival at housing.")

    add_heading(doc, "Step 2: Emergency Evacuation Execution", level=2)
    add_bullet(doc, "1. Alarm Activation: ", "Upon sounding of the building fire alarm, gas alarm, or evacuation order, all personnel must immediately evacuate via the nearest safe emergency exit.")
    add_bullet(doc, "2. No Delay: ", "Personnel shall proceed directly to the designated physical Muster Point. Do not stop to text, collect personal items, or search for co-workers inside the building.")

    add_heading(doc, "Step 3: Headcount & Group Chat Check-In", level=2)
    add_bullet(doc, "1. Initiate Roll Call: ", "Once safely positioned at the Muster Point, the Muster Coordinator sends the standard emergency check-in prompt to the group chat:\n\"EMERGENCY EVACUATION IN PROGRESS. Reply 'SAFE' and your current physical location immediately.\"")
    add_bullet(doc, "2. Member Confirmation: ", "Every group member must reply immediately confirming their safety and presence at the muster point.")
    add_bullet(doc, "3. Verbal / Dead Battery Contingency: ", "If a member's mobile device is dead, lost, or lacking signal, they must physically report to the Muster Coordinator at the muster point so they can be manually checked off.")

    add_heading(doc, "Step 4: Escalation & Reporting Unaccounted Personnel", level=2)
    add_bullet(doc, "1. Time Window (T + 10 Minutes): ", "If any individual who swiped into housing has not checked in via text or in person within ten (10) minutes of the alarm sounding, they are formally declared UNACCOUNTED FOR.")
    add_bullet(doc, "2. Immediate Emergency Notification: ", "The Muster Coordinator must immediately approach the Fire Department Incident Commander or Security Lead at the command post/muster point and report:\n  • Full Name of missing person(s)\n  • Room / Unit Number\n  • Last known location or last time seen\n  • Mobile phone number")

    add_heading(doc, "Step 5: Muster Discipline & All-Clear / Relocation", level=2)
    add_bullet(doc, "1. Hold Position: ", "All group members must remain clustered together at the muster area under the oversight of their Muster Coordinator.")
    add_bullet(doc, "2. Worsening Conditions: ", "If smoke, wind, or hazard plumes threaten the primary muster area, Emergency Services or Security will issue orders to relocate to a Secondary Muster Point. The Muster Coordinator must maintain group cohesion during relocation and re-verify headcount upon arrival.")
    add_bullet(doc, "3. Return to Housing: ", "No personnel may leave the muster area or return to their rooms until the Incident Commander / Fire Department issues a formal, verbal \"ALL-CLEAR.\"")

    # 6. Critical Warnings
    add_heading(doc, "6. CRITICAL LIFE SAFETY WARNINGS", level=1)

    add_callout_box(
        doc,
        "ABSOLUTE PROHIBITION ON RE-ENTRY & SEARCHES",
        "Under NO CIRCUMSTANCES shall any Muster Coordinator, Supervisor, or group member re-enter an evacuated building to search for unaccounted individuals or retrieve belongings.\n\nSearch and Rescue operations inside hazardous structures are strictly reserved for trained, fully equipped Fire Department / Emergency Response personnel wearing Self-Contained Breathing Apparatus (SCBA). Re-entering an evacuated structure turns a rescue into a double fatality.",
        bg_color="F8D7DA",
        border_color="C00000"
    )

    add_callout_box(
        doc,
        "SWIPE-IN ACCURACY IS A LIFE OR DEATH CONTROL",
        "Electronic swipe logs are pulled by Emergency Services during a major incident to determine who is inside the structure.\n\n• If you swipe in and leave without swiping out, firefighters will risk their lives searching burning or toxic rooms for you.\n• If you enter without swiping in and collapse inside, rescuers will not know to look for you.\n\nStrict swipe compliance is mandatory under this interim procedure.",
        bg_color="FFF3CD",
        border_color="856404"
    )

    add_callout_box(
        doc,
        "CELLULAR CONNECTIVITY & BATTERY MITIGATION",
        "Because this interim procedure relies on mobile devices:\n• Keep phones charged while in residence.\n• If cellular data networks become congested during a major site incident, standard SMS (text message) often transmits successfully when mobile internet/apps fail. If data fails, fall back to direct SMS or physical verbal check-in at the muster sign.",
        bg_color="D1ECF1",
        border_color="0C5460"
    )

    # 7. Audit & Review
    add_heading(doc, "7. AUDIT & REVIEW", level=1)
    add_body(doc, "As an interim procedure, this document shall be reviewed by EH&S and Site Management every thirty (30) days to evaluate text group chat reliability, Muster Coordinator effectiveness during drills, and readiness for transition to a permanent automated tracking system.")

    # Save document
    output_path = "Interim_Housing_Evacuation_Headcount_Procedure.docx"
    doc.save(output_path)
    print(f"Successfully generated {output_path}")

if __name__ == "__main__":
    create_document()
