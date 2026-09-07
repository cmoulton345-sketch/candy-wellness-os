"""
Part 3: Assemble the Critical Risk Inventory Excel workbook from JSON data files.
Professional formatting with color-coded risk levels, frozen panes, filters, and column widths.
"""
import json, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

base = os.path.dirname(os.path.abspath(__file__))

# Load data
with open(os.path.join(base, "cri_data_part1.json")) as f:
    data1 = json.load(f)
with open(os.path.join(base, "cri_data_part2.json")) as f:
    data2 = json.load(f)

all_data = data1 + data2

# Column config
columns = [
    ("ID", 6),
    ("Plant Area / System", 28),
    ("Critical Task", 40),
    ("Role", 22),
    ("Operational Phase", 16),
    ("Hazards / Threats", 50),
    ("Potential Consequence", 40),
    ("Uncontrolled Risk Rating", 16),
    ("Critical Controls & Barriers", 55),
    ("Required Competencies / Skills", 45),
    ("Training & Verification Method", 50),
    ("Residual Risk (with controls)", 16),
    ("Regulatory / Standard Reference", 40),
]

field_keys = ["Area","Task","Role","Phase","Hazards","Consequence","Pre-Risk","Controls","Competency","Training","Residual","Reg Ref"]

# Risk color mapping
risk_colors = {
    "Critical": PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid"),
    "High": PatternFill(start_color="FF8C00", end_color="FF8C00", fill_type="solid"),
    "Medium": PatternFill(start_color="FFD700", end_color="FFD700", fill_type="solid"),
    "Low": PatternFill(start_color="32CD32", end_color="32CD32", fill_type="solid"),
}
risk_fonts = {
    "Critical": Font(bold=True, color="FFFFFF", size=10),
    "High": Font(bold=True, color="000000", size=10),
    "Medium": Font(bold=True, color="000000", size=10),
    "Low": Font(bold=True, color="FFFFFF", size=10),
}

# Create workbook
wb = Workbook()

# ============ SHEET 1: CRITICAL RISK INVENTORY ============
ws = wb.active
ws.title = "Critical Risk Inventory"

# Title row
ws.merge_cells("A1:M1")
title_cell = ws["A1"]
title_cell.value = "CRITICAL RISK INVENTORY — LNG Facility (LNG Profile)"
title_cell.font = Font(bold=True, size=16, color="FFFFFF")
title_cell.fill = PatternFill(start_color="1F3864", end_color="1F3864", fill_type="solid")
title_cell.alignment = Alignment(horizontal="center", vertical="center")
ws.row_dimensions[1].height = 40

# Subtitle row
ws.merge_cells("A2:M2")
sub_cell = ws["A2"]
sub_cell.value = "Operators & Maintenance Technicians | Normal Ops · Startup · Shutdown · Emergency | Marine Loading · Remote Site · Wildlife"
sub_cell.font = Font(italic=True, size=11, color="FFFFFF")
sub_cell.fill = PatternFill(start_color="2E5090", end_color="2E5090", fill_type="solid")
sub_cell.alignment = Alignment(horizontal="center", vertical="center")
ws.row_dimensions[2].height = 25

# Header row (row 3)
header_fill = PatternFill(start_color="1F3864", end_color="1F3864", fill_type="solid")
header_font = Font(bold=True, size=10, color="FFFFFF")
header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
thin_border = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"), bottom=Side(style="thin")
)

for col_idx, (col_name, col_width) in enumerate(columns, 1):
    cell = ws.cell(row=3, column=col_idx, value=col_name)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = header_align
    cell.border = thin_border
    ws.column_dimensions[get_column_letter(col_idx)].width = col_width

ws.row_dimensions[3].height = 35

# Data rows
data_font = Font(size=10)
data_align = Alignment(vertical="top", wrap_text=True)

for row_idx, record in enumerate(all_data, 4):
    # ID column
    cell = ws.cell(row=row_idx, column=1, value=f"CRI-{row_idx - 3:03d}")
    cell.font = Font(size=10, bold=True)
    cell.alignment = Alignment(horizontal="center", vertical="top")
    cell.border = thin_border

    # Data columns
    for col_idx, key in enumerate(field_keys, 2):
        val = record.get(key, "")
        cell = ws.cell(row=row_idx, column=col_idx, value=val)
        cell.font = data_font
        cell.alignment = data_align
        cell.border = thin_border

        # Color-code risk columns
        if key in ("Pre-Risk", "Residual") and val in risk_colors:
            cell.fill = risk_colors[val]
            cell.font = risk_fonts[val]
            cell.alignment = Alignment(horizontal="center", vertical="top", wrap_text=True)

    ws.row_dimensions[row_idx].height = 90

# Freeze panes and auto-filter
ws.freeze_panes = "A4"
ws.auto_filter.ref = f"A3:M{3 + len(all_data)}"

# ============ SHEET 2: RISK MATRIX ============
ws2 = wb.create_sheet("Risk Matrix Legend")
ws2.merge_cells("A1:E1")
ws2["A1"].value = "RISK MATRIX — Likelihood × Consequence"
ws2["A1"].font = Font(bold=True, size=14, color="FFFFFF")
ws2["A1"].fill = PatternFill(start_color="1F3864", end_color="1F3864", fill_type="solid")

headers2 = ["", "Catastrophic", "Major", "Moderate", "Minor"]
likelihood = ["Almost Certain", "Likely", "Possible", "Unlikely", "Rare"]
matrix = [
    ["Critical","Critical","High","High"],
    ["Critical","High","High","Medium"],
    ["High","High","Medium","Medium"],
    ["High","Medium","Medium","Low"],
    ["Medium","Medium","Low","Low"],
]

for c, h in enumerate(headers2, 1):
    cell = ws2.cell(row=3, column=c, value=h)
    cell.font = Font(bold=True, size=10, color="FFFFFF")
    cell.fill = PatternFill(start_color="1F3864", end_color="1F3864", fill_type="solid")
    cell.alignment = Alignment(horizontal="center")
    cell.border = thin_border
    ws2.column_dimensions[get_column_letter(c)].width = 18

for r, (lh, row_vals) in enumerate(zip(likelihood, matrix), 4):
    cell = ws2.cell(row=r, column=1, value=lh)
    cell.font = Font(bold=True, size=10)
    cell.fill = PatternFill(start_color="D9E2F3", end_color="D9E2F3", fill_type="solid")
    cell.border = thin_border
    for c, val in enumerate(row_vals, 2):
        cell = ws2.cell(row=r, column=c, value=val)
        cell.fill = risk_colors.get(val, PatternFill())
        cell.font = risk_fonts.get(val, Font(size=10))
        cell.alignment = Alignment(horizontal="center")
        cell.border = thin_border

# Legend
ws2.cell(row=10, column=1, value="Risk Rating Definitions:").font = Font(bold=True, size=11)
defs = [
    ("Critical", "Immediate risk to life. Task requires highest level of control. Stop work authority applies. Senior management approval required."),
    ("High", "Significant risk. Engineering and administrative controls mandatory. Competency verification required before independent task performance."),
    ("Medium", "Moderate risk. Standard controls apply. Supervision required for new/inexperienced personnel."),
    ("Low", "Risk managed through routine procedures and general awareness training."),
]
for i, (rating, desc) in enumerate(defs):
    r = 11 + i
    cell = ws2.cell(row=r, column=1, value=rating)
    cell.fill = risk_colors[rating]
    cell.font = risk_fonts[rating]
    cell.border = thin_border
    ws2.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5)
    cell2 = ws2.cell(row=r, column=2, value=desc)
    cell2.font = Font(size=10)
    cell2.alignment = Alignment(wrap_text=True)
    cell2.border = thin_border

# ============ SHEET 3: COMPETENCY ASSURANCE PLAN ============
ws3 = wb.create_sheet("Competency Assurance Plan")
ws3.merge_cells("A1:F1")
ws3["A1"].value = "COMPETENCY ASSURANCE PLAN — Operational Readiness Framework"
ws3["A1"].font = Font(bold=True, size=14, color="FFFFFF")
ws3["A1"].fill = PatternFill(start_color="1F3864", end_color="1F3864", fill_type="solid")

cap_headers = ["Phase", "Timeline", "Activity", "Responsible", "Deliverable", "Verification Method"]
cap_widths = [20, 18, 45, 22, 35, 35]
for c, (h, w) in enumerate(zip(cap_headers, cap_widths), 1):
    cell = ws3.cell(row=3, column=c, value=h)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = header_align
    cell.border = thin_border
    ws3.column_dimensions[get_column_letter(c)].width = w

cap_data = [
    ["Phase 1: Foundation", "T-18 to T-12 months", "Develop role-specific competency profiles from CRI", "HSE Manager + Ops Manager", "Competency profile document per role", "Management review and sign-off"],
    ["Phase 1: Foundation", "T-18 to T-12 months", "Build training curriculum aligned to CRI hazards", "Training Coordinator", "Training matrix mapped to each CRI item", "Gap analysis against CRI"],
    ["Phase 1: Foundation", "T-18 to T-12 months", "Procure Operator Training Simulator (OTS)", "Engineering + Ops", "Commissioned OTS with plant-specific model", "Factory Acceptance Test (FAT)"],
    ["Phase 2: Hiring & Initial Training", "T-12 to T-6 months", "Recruit experienced LNG operators and maintenance techs", "HR + Ops Manager", "Staffing plan filled to minimum levels", "Resume verification + reference checks"],
    ["Phase 2: Hiring & Initial Training", "T-12 to T-6 months", "Deliver LNG fundamentals classroom training", "Training Coordinator", "Completion certificates for all staff", "Written exam (min 80% pass)"],
    ["Phase 2: Hiring & Initial Training", "T-12 to T-6 months", "Complete mandatory safety certifications", "HSE Manager", "H2S Alive, Confined Space, Fall Protection, LOTO, WHMIS, First Aid, Bear Aware certs on file", "Certificate verification + practical demo"],
    ["Phase 3: Simulator & Practical", "T-6 to T-3 months", "OTS training: Startup, shutdown, emergency scenarios", "Lead Operator + Training Coord", "Min 5 startup + 5 shutdown simulations per operator; all emergency scenarios completed", "OTS performance assessment (pass/fail criteria)"],
    ["Phase 3: Simulator & Practical", "T-6 to T-3 months", "Hands-on equipment familiarization (site walk-downs)", "Ops Supervisors", "Equipment location knowledge verified", "Practical: Identify equipment from P&ID in field"],
    ["Phase 3: Simulator & Practical", "T-6 to T-3 months", "Emergency response drills (ESD, fire, H2S, spill, man overboard)", "HSE + ERT Lead", "Drill completion records with corrective actions", "Drill evaluation + debrief action items closed"],
    ["Phase 4: Commissioning OJT", "T-3 to T-0 (First Gas)", "On-the-job training during commissioning activities", "Senior Operators / Supervisors", "OJT logbook signed off for each critical task", "Supervisor sign-off + independent assessment"],
    ["Phase 4: Commissioning OJT", "T-3 to T-0 (First Gas)", "Pre-Startup Safety Review (PSSR) participation", "All Ops + Maintenance", "PSSR checklist completed and signed", "PSSR audit by independent reviewer"],
    ["Phase 4: Commissioning OJT", "T-3 to T-0 (First Gas)", "Marine loading arm connection/disconnection practical", "Marine Ops Lead", "Min 10 supervised connections per operator", "Practical assessment with Marine Ops Lead sign-off"],
    ["Phase 5: Ongoing Assurance", "Post-Startup (Continuous)", "Annual competency re-assessment for all critical tasks", "HSE + Ops Managers", "Updated competency matrix with gap closure plans", "Annual assessment + refresher training records"],
    ["Phase 5: Ongoing Assurance", "Post-Startup (Continuous)", "Emergency drill program (quarterly fire, semi-annual H2S, annual full-scale)", "HSE Manager + ERT Lead", "Drill schedule adherence >95%; corrective actions closed within 30 days", "Drill reports + regulatory inspection readiness"],
    ["Phase 5: Ongoing Assurance", "Post-Startup (Continuous)", "Management of Change (MOC) - competency impact review", "All Managers", "MOC register with training impact assessment for each change", "MOC audit (random sample quarterly)"],
    ["Phase 5: Ongoing Assurance", "Post-Startup (Continuous)", "Lessons learned integration from incidents and near-misses", "HSE Manager", "Updated SOPs + CRI + training materials", "Incident review board sign-off"],
]

for r, row_data in enumerate(cap_data, 4):
    for c, val in enumerate(row_data, 1):
        cell = ws3.cell(row=r, column=c, value=val)
        cell.font = Font(size=10)
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        cell.border = thin_border
    ws3.row_dimensions[r].height = 50

ws3.freeze_panes = "A4"

# ============ SHEET 4: OPEN QUESTIONS ============
ws4 = wb.create_sheet("Open Questions for Joe")
ws4.merge_cells("A1:C1")
ws4["A1"].value = "OPEN QUESTIONS — Required to Finalize CRI for Operational Readiness"
ws4["A1"].font = Font(bold=True, size=14, color="FFFFFF")
ws4["A1"].fill = PatternFill(start_color="1F3864", end_color="1F3864", fill_type="solid")

q_headers = ["#", "Question", "Impact on CRI"]
q_widths = [5, 65, 50]
for c, (h, w) in enumerate(zip(q_headers, q_widths), 1):
    cell = ws4.cell(row=3, column=c, value=h)
    cell.font = header_font; cell.fill = header_fill; cell.alignment = header_align; cell.border = thin_border
    ws4.column_dimensions[get_column_letter(c)].width = w

questions = [
    ("Is LNG using a single train or multiple liquefaction trains?", "Impacts redundancy assumptions and simultaneous operations (SIMOPS) risk entries"),
    ("What is the shift rotation pattern (e.g., 14/14, 7/7, 4x4)?", "Affects fatigue management controls and training scheduling in the Competency Plan"),
    ("Are there any SIMOPS planned (e.g., construction ongoing during commissioning)?", "Would require additional CRI entries for interface risks between construction and operations crews"),
    ("What specific ERT (Emergency Response Team) model is planned — dedicated on-site or mutual aid?", "Determines rescue capability assumptions in confined space, H2S, fire, and marine entries"),
    ("Is there a helipad for medevac, or is marine ambulance the primary evacuation route?", "Affects remote site medical emergency response controls and competency requirements"),
    ("What DCS/SCADA platform is being used?", "Determines OTS simulator procurement and control room operator training specifics"),
    ("Are there any indigenous cultural protocols or Squamish Nation requirements for site access/training?", "May require additional cultural safety training entries in the Competency Assurance Plan"),
    ("What is the flare system capacity and any community noise/emission restrictions?", "Affects startup/shutdown procedure constraints and emergency depressurization controls"),
    ("Will maintenance be in-house or contract-based (or hybrid)?", "Impacts contractor management entries, Competency Assurance scope, and LOTO multi-employer coordination"),
    ("Are there any process safety KPIs already defined by the operator (e.g., Tier 1/Tier 2 LOPC tracking)?", "Determines if additional monitoring/reporting entries are needed in ongoing assurance phase"),
]

for r, (q, impact) in enumerate(questions, 4):
    ws4.cell(row=r, column=1, value=r-3).font = Font(size=10, bold=True)
    ws4.cell(row=r, column=1).alignment = Alignment(horizontal="center", vertical="top")
    ws4.cell(row=r, column=1).border = thin_border
    c2 = ws4.cell(row=r, column=2, value=q)
    c2.font = Font(size=10); c2.alignment = Alignment(wrap_text=True, vertical="top"); c2.border = thin_border
    c3 = ws4.cell(row=r, column=3, value=impact)
    c3.font = Font(size=10); c3.alignment = Alignment(wrap_text=True, vertical="top"); c3.border = thin_border
    ws4.row_dimensions[r].height = 40

ws4.freeze_panes = "A4"

# ============ SAVE ============
out_path = os.path.join(base, "LNG_LNG_Critical_Risk_Inventory.xlsx")
wb.save(out_path)
print(f"Excel workbook saved: {out_path}")
print(f"Total CRI entries: {len(all_data)}")
print(f"Sheets: {wb.sheetnames}")
