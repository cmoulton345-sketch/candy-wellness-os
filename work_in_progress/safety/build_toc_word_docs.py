import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

def set_cell_background(cell, fill_color):
    """Sets background color of a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Sets cell padding in dxa."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_header_footer(doc, doc_title):
    section = doc.sections[0]
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    
    # Header
    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hrun = hp.add_run(f"BC LNG Facility | {doc_title}")
    hrun.font.name = 'Calibri'
    hrun.font.size = Pt(8.5)
    hrun.font.color.rgb = RGBColor(128, 128, 128)
    
    # Footer
    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    frun = fp.add_run("CONTROLLED DOCUMENT - UNCONTROLLED WHEN PRINTED")
    frun.font.name = 'Calibri'
    frun.font.size = Pt(8.5)
    frun.font.color.rgb = RGBColor(128, 128, 128)

def create_doc_control_table(doc, doc_no, title, version="1.0"):
    table = doc.add_table(rows=5, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    col_widths = [Inches(2.2), Inches(4.3)]
    
    data = [
        ("Document Number:", doc_no),
        ("Document Title:", title),
        ("Version & Status:", f"{version} (Master Architecture - Work In Progress)"),
        ("Jurisdiction & Governing Codes:", "BC Energy Regulator (BCER), WorkSafeBC OHSR, TSBC, CSA Z276, ISO 45001"),
        ("Target Audience:", "Owner Leadership, Operations & Commissioning, EPC Contractor, Regulators")
    ]
    
    for i, (label, val) in enumerate(data):
        row = table.rows[i]
        
        # Label cell
        cell_lbl = row.cells[0]
        cell_lbl.width = col_widths[0]
        set_cell_background(cell_lbl, "1B365D") # Navy
        set_cell_margins(cell_lbl, top=80, bottom=80, left=120, right=120)
        p = cell_lbl.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(label)
        r.font.name = 'Calibri'
        r.font.bold = True
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(255, 255, 255)
        
        # Value cell
        cell_val = row.cells[1]
        cell_val.width = col_widths[1]
        set_cell_background(cell_val, "F4F6F9") # Soft light gray
        set_cell_margins(cell_val, top=80, bottom=80, left=120, right=120)
        p2 = cell_val.paragraphs[0]
        p2.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r2 = p2.add_run(val)
        r2.font.name = 'Calibri'
        r2.font.size = Pt(9.5)
        r2.font.color.rgb = RGBColor(30, 30, 30)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

def build_osmm_docx(output_path):
    doc = docx.Document()
    add_header_footer(doc, "Operations Safety Management Manual (OSMM) TOC")
    
    # Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r_title = p_title.add_run("OPERATIONS SAFETY MANAGEMENT MANUAL (OSMM)")
    r_title.font.name = 'Calibri'
    r_title.font.size = Pt(22)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(27, 54, 93) # Navy #1B365D
    p_title.paragraph_format.space_after = Pt(2)
    
    p_sub = doc.add_paragraph()
    r_sub = p_sub.add_run("Master Table of Contents & ISO 45001 / BCER Compliance Structure")
    r_sub.font.name = 'Calibri'
    r_sub.font.size = Pt(13)
    r_sub.font.italic = True
    r_sub.font.color.rgb = RGBColor(0, 90, 156) # Teal #005A9C
    p_sub.paragraph_format.space_after = Pt(16)
    
    # Doc control table
    create_doc_control_table(doc, "OSMM-MAN-001", "Operations Safety Management Manual", "1.0")
    
    # Content structure
    sections = [
        ("SECTION 1.0: CONTEXT OF THE ORGANIZATION & REGULATORY SCOPE (ISO 45001 Clause 4)", [
            ("1.1 Facility Overview & Physical Boundaries (ISO 45001 Cl. 4.1)", [
                "1.1.1 Asset Description (Liquefaction, Cryogenic Storage, Marine Terminal, Utilities)",
                "1.1.2 Battery Limits & Geographical Footprint",
                "1.1.3 Operating Envelopes & Design Limits"
            ]),
            ("1.2 Statutory, Regulatory & Legal Compliance Framework (ISO 45001 Cl. 4.2 & 6.1.3)", [
                "1.2.1 BC Energy Regulator (BCER) Permit & LNG Facility Regulation Register",
                "1.2.2 WorkSafeBC Occupational Health and Safety Regulation (OHSR) Compliance Register",
                "1.2.3 Technical Safety BC (TSBC) Safety Standards Act Register",
                "1.2.4 Environmental Assessment Office (EAO) & Federal Impact Assessment (IAA) Conditions",
                "1.2.5 Applicable Canadian Standards Association (CSA) Codes (CSA Z276, CSA Z246.1, CSA Z246.2)"
            ]),
            ("1.3 Scope of the Safety Management System (SMS) (ISO 45001 Cl. 4.3)", [
                "1.3.1 Inclusions, Exclusions, and Lifecycle Application (Operations through Decommissioning)",
                "1.3.2 Integration with Environmental (ISO 14001) and Asset Management (ISO 55001) Systems"
            ]),
            ("1.4 Interested Parties & Stakeholder Expectations (ISO 45001 Cl. 4.2)", [
                "1.4.1 Worker & Union Representatives",
                "1.4.2 Regulatory Bodies (BCER, WorkSafeBC, TSBC, Transport Canada)",
                "1.4.3 First Nations & Indigenous Community Consultation Requirements",
                "1.4.4 Local Municipalities & Emergency Services"
            ])
        ]),
        ("SECTION 2.0: LEADERSHIP, COMMITMENT & WORKER PARTICIPATION (ISO 45001 Clause 5)", [
            ("2.1 Executive Leadership & Safety Policy (ISO 45001 Cl. 5.1 & 5.2)", [
                "2.1.1 Health, Safety, Environment & Process Safety Policy Statement",
                "2.1.2 Executive Commitment to Hazard Elimination and Risk Reduction"
            ]),
            ("2.2 Governance, Roles, Accountabilities & Authorities (ISO 45001 Cl. 5.3)", [
                "2.2.1 Organizational Structure & Safety Leadership Chain",
                "2.2.2 Key Roles & Statutory Accountabilities (Plant Manager, Process Safety Manager, Operations Lead)",
                "2.2.3 WorkSafeBC Prime Contractor Designation & Duty Delegation (Section 118 WCA)"
            ]),
            ("2.3 Worker Consultation, Participation & Empowerment (ISO 45001 Cl. 5.4)", [
                "2.3.1 Joint Health and Safety Committee (JHSC) Structure & Mandate (WorkSafeBC Part 3)",
                "2.3.2 Worker Representation in Hazard Identification & Risk Assessment"
            ]),
            ("2.4 Rights & Protection of Personnel", [
                "2.4.1 Right to Refuse Unsafe Work Protocol (WorkSafeBC OHSR 3.12)",
                "2.4.2 Stop Work Authority (SWA) Policy & Protection against Retaliation"
            ])
        ]),
        ("SECTION 3.0: PLANNING — HAZARD IDENTIFICATION, RISK ASSESSMENT & LEGAL COMPLIANCE (ISO 45001 Clause 6)", [
            ("3.1 Hazard Identification & Risk Assessment (HIRA) Program (ISO 45001 Cl. 6.1.1 & 6.1.2)", [
                "3.1.1 Occupational Hazard Identification Methodology",
                "3.1.2 Risk Assessment Matrix & Tolerability Criteria (ALARP Principles)"
            ]),
            ("3.2 Process Safety Management (PSM) Framework (CCPS / CSA Z276)", [
                "3.2.1 Major Accident Hazard (MAH) Identification",
                "3.2.2 Hazard & Operability Studies (HAZOP) and Layer of Protection Analysis (LOPA)",
                "3.2.3 Bowtie Risk Analysis & Safety Critical Element (SCE) Linkage"
            ]),
            ("3.3 Legal & Regulatory Requirements Management (ISO 45001 Cl. 6.1.3)", [
                "3.3.1 Compliance Identification, Monitoring, and Updating Procedure",
                "3.3.2 Regulatory Reporting & Notification Matrix"
            ]),
            ("3.4 SMS Objectives, Targets & Action Planning (ISO 45001 Cl. 6.2)", [
                "3.4.1 Annual Health, Safety & Process Safety Objectives",
                "3.4.2 Implementation Programs, Timelines, and Resource Allocation"
            ])
        ]),
        ("SECTION 4.0: SUPPORT SYSTEMS — COMPETENCE, COMMUNICATION & DOCUMENTATION (ISO 45001 Clause 7)", [
            ("4.1 Resource & Infrastructure Management (ISO 45001 Cl. 7.1)", [
                "4.1.1 Staffing Levels & Competent Person Allocations",
                "4.1.2 Safety Facilities, Equipment, and Personal Protective Equipment (PPE) Standards"
            ]),
            ("4.2 Training, Competency & Qualifications Assurance (ISO 45001 Cl. 7.2)", [
                "4.2.1 Training Needs Analysis (TNA) & Role Competency Matrix",
                "4.2.2 Operator Certification, High-Pressure Boiler/Refrigeration Qualifications (TSBC)",
                "4.2.3 Competency Verification & Re-Certification Intervals"
            ]),
            ("4.3 Safety Awareness & Culture Programs (ISO 45001 Cl. 7.3)", [
                "4.3.1 Mandatory Site Orientation Programs (General Site, Visitor, Contractor)",
                "4.3.2 Safety Moments, Toolbox Talks & Campaign Frameworks"
            ]),
            ("4.4 Communication, Consultation & Reporting Protocols (ISO 45001 Cl. 7.4)", [
                "4.4.1 Internal Safety Communication Channels & Shift Handover Standards",
                "4.4.2 External & Regulatory Communication Protocols"
            ]),
            ("4.5 Documented Information & Records Control (ISO 45001 Cl. 7.5)", [
                "4.5.1 SMS Document Hierarchy, Formatting, and Approval Workflows",
                "4.5.2 Control of Records, Archiving, and Statutory Retention Periods"
            ])
        ]),
        ("SECTION 5.0: OPERATIONAL PLANNING & CONTROL (ISO 45001 Clause 8.1)", [
            ("5.1 Operational Planning & Control Governance (ISO 45001 Cl. 8.1.1)", [
                "5.1.1 Standard Operating Procedures (SOPs) & Operating Manuals",
                "5.1.2 Operating Boundaries & Alarm Management Protocols"
            ]),
            ("5.2 Safe Work Permitting System (PTW) (ISO 45001 Cl. 8.1.1)", [
                "5.2.1 Baseline Permit to Work (PTW) Architecture & Roles (Issuing / Performing Authorities)",
                "5.2.2 Permit Categories (Cold Work, Hot Work Class A/B, Confined Space, Critical Lift)"
            ]),
            ("5.3 Energy Isolation & Lockout/Tagout (LOTO) System (WorkSafeBC OHSR Part 10)", [
                "5.3.1 Zero Energy State Verification & Isolation Standards (Mechanical, Electrical, Process)",
                "5.3.2 Blind/Blank Registry and High-Pressure Isolation Protocols"
            ]),
            ("5.4 High-Hazard Safe Work Practices", [
                "5.4.1 Confined Space Entry & Atmospheric Testing (WorkSafeBC Part 9)",
                "5.4.2 Working at Heights & Fall Protection Governance (WorkSafeBC Part 11)",
                "5.4.3 Excavation, Trenching, and Ground Disturbance (WorkSafeBC Part 20)"
            ]),
            ("5.5 Management of Change (MOC) System (ISO 45001 Cl. 8.1.3)", [
                "5.5.1 Technical MOC (Process, Piping, Instrumentation, Control Logic)",
                "5.5.2 Organizational & Procedural MOC Protocols",
                "5.5.3 Pre-Startup Safety Review (PSSR) Standards"
            ]),
            ("5.6 Asset Integrity & Facility Integrity Management Plan (FIMP) (ISO 55001 / CSA Z276)", [
                "5.6.1 Pressure Equipment, Piping & Cryogenic Tank Inspection (TSBC / API Standards)",
                "5.6.2 Safety Critical Elements (SCE) Performance Standards & Testing Intervals",
                "5.6.3 Safety System Impairment & Bypass Authorization Controls"
            ]),
            ("5.7 Multi-Employer Workplace & Contractor Safety Management (ISO 45001 Cl. 8.1.4)", [
                "5.7.1 Prime Contractor Site Safety Coordination Rules (WorkSafeBC Sec 118)",
                "5.7.2 Contractor Pre-Qualification, Auditing, and Safety Performance Management"
            ])
        ]),
        ("SECTION 6.0: EMERGENCY PREPAREDNESS & RESPONSE (ISO 45001 Clause 8.2 / CSA Z246.2)", [
            ("6.1 Emergency Response Plan (ERP) Architecture (CSA Z246.2 & BCER)", [
                "6.1.1 Incident Command System (ICS) Structure & Command Posts",
                "6.1.2 Hazard-Specific Emergency Procedures (LNG Release, Cryogenic Burn, BLEVE, Structural Fire)"
            ]),
            ("6.2 Emergency Resources & Medical Integration", [
                "6.2.1 Dedicated 3rd-Party Emergency/Medical Services Provider Protocol",
                "6.2.2 Firefighting Systems, Foam Systems, and Emergency Isolation Valve (EIV) Operations"
            ]),
            ("6.3 Testing, Drills & Readiness Assurance", [
                "6.3.1 Mandatory Drill Schedule (Tabletop, Functional, Full-Scale Exercises)",
                "6.3.2 ERP Training & Emergency Response Team (ERT) Competency"
            ]),
            ("6.4 External Coordination & Mutual Aid", [
                "6.4.1 Municipal, Regional, and First Nations Emergency Services Alignment",
                "6.4.2 Regulatory Emergency Reporting (BCER Emergency Management Regulation)"
            ])
        ]),
        ("SECTION 7.0: PERFORMANCE EVALUATION & AUDIT (ISO 45001 Clause 9)", [
            ("7.1 Performance Monitoring, Measurement & Analysis (ISO 45001 Cl. 9.1.1)", [
                "7.1.1 Leading Indicators (Audits, Near-Misses, Hazard Cards, MOC Closures)",
                "7.1.2 Lagging Indicators (TRIF, DART, Loss of Primary Containment - LOPC Events)"
            ]),
            ("7.2 Evaluation of Compliance (ISO 45001 Cl. 9.1.2)", [
                "7.2.1 Periodic Statutory Compliance Audits (BCER, WorkSafeBC, TSBC)",
                "7.2.2 Environmental & EAO Certificate Compliance Reviews"
            ]),
            ("7.3 Internal Audit Program (ISO 45001 Cl. 9.2)", [
                "7.3.1 Annual SMS & Process Safety Audit Schedule",
                "7.3.2 Auditor Competency, Reporting, and Findings Resolution"
            ]),
            ("7.4 Executive Management Review (ISO 45001 Cl. 9.3)", [
                "7.4.1 Annual Management Review Inputs & Governance Meetings",
                "7.4.2 Strategic Directives & System Resource Allocation"
            ])
        ]),
        ("SECTION 8.0: INCIDENT MANAGEMENT & CONTINUAL IMPROVEMENT (ISO 45001 Clause 10)", [
            ("8.1 Incident, Injury & Near-Miss Reporting (ISO 45001 Cl. 10.2)", [
                "8.1.1 Worker Incident Reporting Workflows & Initial Triage",
                "8.1.2 Incident Severity Rating Matrix & Investigation Scopes"
            ]),
            ("8.2 Statutory Incident Notifications", [
                "8.2.1 BCER Mandatory Incident Reporting (Section 11 LNG Facility Regulation)",
                "8.2.2 WorkSafeBC Immediate Notification (Serious Injury / Fatality / Critical Breakdown)",
                "8.2.3 Transportation Safety Board (TSB) & TSBC Incident Notification"
            ]),
            ("8.3 Root Cause Analysis (RCA) & Corrective Actions (CAPA) (ISO 45001 Cl. 10.2)", [
                "8.3.1 Root Cause Investigation Methodologies (TapRooT, 5-Why, Cause Tree)",
                "8.3.2 Corrective & Preventive Action (CAPA) Tracking, Verification & Effectiveness Reviews"
            ]),
            ("8.4 Continual Improvement Framework (ISO 45001 Cl. 10.3)", [
                "8.4.1 Lessons Learned Integration into SOPs & Training",
                "8.4.2 Systematic SMS Evolution & Policy Updates"
            ])
        ])
    ]
    
    render_sections(doc, sections)
    doc.save(output_path)
    print(f"Saved: {output_path}")

def build_tsmp_docx(output_path):
    doc = docx.Document()
    add_header_footer(doc, "Transitional Safety Management Plan (TSMP) TOC")
    
    # Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r_title = p_title.add_run("TRANSITIONAL SAFETY MANAGEMENT PLAN (TSMP)")
    r_title.font.name = 'Calibri'
    r_title.font.size = Pt(22)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(27, 54, 93) # Navy #1B365D
    p_title.paragraph_format.space_after = Pt(2)
    
    p_sub = doc.add_paragraph()
    r_sub = p_sub.add_run("Master Table of Contents & Commissioning / Handover Execution Plan")
    r_sub.font.name = 'Calibri'
    r_sub.font.size = Pt(13)
    r_sub.font.italic = True
    r_sub.font.color.rgb = RGBColor(0, 90, 156) # Teal #005A9C
    p_sub.paragraph_format.space_after = Pt(16)
    
    # Doc control table
    create_doc_control_table(doc, "TSMP-PLN-001", "Transitional Safety Management Plan", "1.0")
    
    sections = [
        ("SECTION 1.0: PURPOSE, SCOPE & TRANSITION MILESTONES", [
            ("1.1 Purpose & Execution Scope", [
                "1.1.1 Governance Bridge between EPC Construction Rules and Owner OSMM",
                "1.1.2 Lifecycle Phases: Mechanical Completion (MC) -> Pre-Commissioning -> Commissioning -> Hydrocarbon Intro -> Handover"
            ]),
            ("1.2 Dispersed / Subsystem Turnover Execution Model", [
                "1.2.1 Non-Linear System Handover Strategy (Early Utility Systems e.g., Instrument Air, Nitrogen, Power)",
                "1.2.2 Interface Control between Energized Subsystems and Adjacent Active Construction Footprints"
            ]),
            ("1.3 Statutory & Regulatory Approvals Gate", [
                "1.3.1 BCER Commissioning Consent & Testing Notifications",
                "1.3.2 Technical Safety BC (TSBC) Operating Permits & Inspection Sign-Offs"
            ])
        ]),
        ("SECTION 2.0: TRANSITIONAL ROLES, RESPONSIBILITIES & AUTHORITY MATRIX", [
            ("2.1 WorkSafeBC Prime Contractor Governance during Transition", [
                "2.1.1 Owner Prime Contractor Overall Site Safety Authority",
                "2.1.2 Owner Safety Oversight Structure in Multi-Employer Environment"
            ]),
            ("2.2 EPC Contractor Oversight & Subcontractor Safety Alignment", [
                "2.2.1 EPC Construction Oversight Responsibilities on Un-Turned-Over Assets",
                "2.2.2 Subcontractor Safety Alignment & Compliance Enforcement"
            ]),
            ("2.3 Owner Operations & Commissioning Team Authority", [
                "2.3.1 Commissioning Manager & Area Isolation Authority Powers",
                "2.3.2 Operations Permit Issuing Authority on Turned-Over Assets"
            ]),
            ("2.4 Interface & Authority Matrix (RACIS Chart)", [
                "2.4.1 System Handover RACIS Matrix (Construction vs. Operations vs. Safety vs. EPC)"
            ])
        ]),
        ("SECTION 3.0: CARE, CUSTODY, AND CONTROL (CCCC) & HANDOVER GATES", [
            ("3.1 System Boundary Demarcation & Red Line Markings", [
                "3.1.1 System Demarcation P&ID Drawings & Color-Coded Piping Maps",
                "3.1.2 Physical Boundary Flags, Barricades, and Signage Criteria"
            ]),
            ("3.2 Joint System Walkdown & Punchlisting Protocol", [
                "3.2.1 Pre-Walkdown Documentation Verification (Hydrotest, NDE, Flush/Dry Records)",
                "3.2.2 Joint Walkdown Team Composition (Owner Engineering, Operations, EPC, Quality)"
            ]),
            ("3.3 Punchlist Categorization & Mandatory Closeout Gates", [
                "3.3.1 Category A Punch Items: Mandatory Closeout Prior to Mechanical Completion / Testing",
                "3.3.2 Category B & C Punch Items: Closure Requirements Prior to Commodity Introduction"
            ]),
            ("3.4 Transfer of Care, Custody, and Control (CCCC) Certificates", [
                "3.4.1 Formal Certificate Execution Sign-Off Workflow",
                "3.4.2 Legal Asset Custody Transfer Logging"
            ])
        ]),
        ("SECTION 4.0: TRANSITIONAL PERMIT TO WORK (PTW) PROTOCOL", [
            ("4.1 Harmonized PTW Architecture", [
                "4.1.1 Retention of EPC Permit Form Factor to Preserve Workforce Familiarity",
                "4.1.2 Transfer of Permitting Authority to Owner Operations for Turned-Over Systems"
            ]),
            ("4.2 Daily Scope-Specific Permitting Execution", [
                "4.2.1 Mandatory Shift to Daily Permits for Turned-Over Systems",
                "4.2.2 Daily Permit Issuance Workflows, Joint Field Risk Assessments & Gas Testing"
            ]),
            ("4.3 Multi-Day 'Blanket Permit' Rules & Risk Screening", [
                "4.3.1 Eligible Low-Risk Work Scopes (Static Inspections, Non-Intrusive Civil/Painting)",
                "4.3.2 Strict Exclusion Criteria & Revocation Triggers"
            ]),
            ("4.4 PTW Interface & Cross-Boundary Permitting Controls", [
                "4.4.1 Work Originating in Construction Zones Impacting Live Systems (Dual Permitting Sign-Off)"
            ])
        ]),
        ("SECTION 5.0: PHYSICAL, VISUAL & COMMUNICATION ISOLATION CONTROLS", [
            ("5.1 Physical Energy Isolation (LOTO) Controls", [
                "5.1.1 Operations Color-Coded Lock System & Key Management",
                "5.1.2 Specialty Positive Isolation: Blinds, Blanks, Double-Block & Bleed Standards"
            ]),
            ("5.2 Visual Boundary Demarcation", [
                "5.2.1 High-Visibility Live System Tags, Color-Coded Blinding Tags, and Danger Signage",
                "5.2.2 Physical Cable & Chain Isolation of Operations-Controlled Valves"
            ]),
            ("5.3 Advance 'Livening Up' Notification Schedule", [
                "5.3.1 Mandatory 7-Day Advance Site-Wide Livening Notification & Map Broadcast",
                "5.3.2 Mandatory 48-Hour Re-Notification & Area Supervisor Alignment",
                "5.3.3 Mandatory 24-Hour Final Warning, Horn/Siren Broadcast & Physical Walkdown"
            ]),
            ("5.4 Lockout/Tagout (LOTO) Custody Transfer Mechanics", [
                "5.4.1 Formal Cutting/Removal of EPC Construction Locks",
                "5.4.2 Installation of Operations Master Locks on Centralized Lock Boxes"
            ])
        ]),
        ("SECTION 6.0: SIMOPS MANAGEMENT & MULTI-LAYERED FIELD ASSURANCE", [
            ("6.1 Dispersed System SIMOPS Hazard Analysis", [
                "6.1.1 SIMOPS Risk Matrix (Construction vs. Energized Utilities vs. Hydrocarbons)",
                "6.1.2 Area-Specific SIMOPS Risk Mitigation Plans"
            ]),
            ("6.2 Multi-Layered Field Assurance & Verification System", [
                "6.2.1 Safety Department Field Auditing & Formal Compliance Reporting",
                "6.2.2 Supervisory Field Verification Tools & Daily Observation Checklists",
                "6.2.3 Mandatory Stage & Action Gate Sign-Off Workflows"
            ])
        ]),
        ("SECTION 7.0: MANDATORY IN-PERSON TRANSITION ORIENTATION & COMPETENCY", [
            ("7.1 Transition Orientation Mandate", [
                "7.1.1 100% Mandatory Attendance for All Personnel On-Site Prior to System Energization",
                "7.1.2 Passing Score Threshold & Badge Re-Authorization Protocol"
            ]),
            ("7.2 In-Person Delivery & Verification Standards", [
                "7.2.1 Classroom In-Person Curriculum (Live System Identification, LOTO Rules, PTW Interface)",
                "7.2.2 Written & Practical Comprehension Testing"
            ]),
            ("7.3 Train-the-Trainer Program & Authorized Instructor Designation", [
                "7.3.1 Instructor Qualification Criteria & Approved Trainer List"
            ])
        ]),
        ("SECTION 8.0: MANAGEMENT OF CHANGE (MOC) ON TURNED-OVER ASSETS", [
            ("8.1 MOC Mandatory Triggers on Turned-Over Systems", [
                "8.1.1 Modification Rules for Pipe, Instrument, Control Logic, or Valve Re-Configuration",
                "8.1.2 Temporary Modifications & Emergency MOC Protocols"
            ]),
            ("8.2 Joint MOC Review Committee", [
                "8.2.1 Committee Composition (Engineering, Operations, Commissioning, Safety)",
                "8.2.2 MOC Risk Assessment, Approval Sign-Off, and Red-Line Drawing Updates"
            ]),
            ("8.3 Field Work Execution & Re-Testing Controls", [])
        ]),
        ("SECTION 9.0: PRE-STARTUP SAFETY REVIEW (PSSR) & COMMODITY INTRODUCTION GATE", [
            ("9.1 Pre-Commissioning & Testing Protocols", [
                "9.1.1 Nitrogen Purging, Dewatering, Drying, and Vacuum Testing Controls",
                "9.1.2 Cryogenic Leak Testing & Tightness Verification Procedures"
            ]),
            ("9.2 Pre-Startup Safety Review (PSSR) Gate Execution", [
                "9.2.1 Multidisciplinary PSSR Team Composition & Field Verification Checklists",
                "9.2.2 PSSR Approval Certificate Sign-Off Workflow"
            ]),
            ("9.3 Strict Punchlist Closure Gate", [
                "9.3.1 Mandatory 100% Closure of Category A, B, and C Punch Items Prior to Hydrocarbon Intro"
            ])
        ]),
        ("SECTION 10.0: PHASED EMERGENCY MANAGEMENT & ICS TRANSITION", [
            ("10.1 Phase 1 Emergency Response Structure (Construction & Pre-Commissioning)", [
                "10.1.1 Owner-Managed 3rd-Party Medical & Emergency Services Provider Role",
                "10.1.2 Site Medical, Muster Point, and Construction Siren Operations"
            ]),
            ("10.2 Phase 2 Emergency Response Structure (Commissioning & Testing)", [
                "10.2.1 Shared Dual Operations & Construction Incident Command System (ICS)",
                "10.2.2 Dual Incident Commander Protocols & Unified Command Post"
            ]),
            ("10.3 Phase 3 Emergency Response Structure (Hydrocarbon Introduction & Onward)", [
                "10.3.1 Transition to Sole Operations Incident Command",
                "10.3.2 EPC & Construction Support Roles during Demobilization"
            ])
        ])
    ]
    
    render_sections(doc, sections)
    doc.save(output_path)
    print(f"Saved: {output_path}")

def render_sections(doc, sections):
    for sec_title, subsecs in sections:
        # Heading 1 (Section)
        p1 = doc.add_paragraph()
        p1.paragraph_format.space_before = Pt(14)
        p1.paragraph_format.space_after = Pt(4)
        p1.paragraph_format.keep_with_next = True
        r1 = p1.add_run(sec_title)
        r1.font.name = 'Calibri'
        r1.font.size = Pt(12.5)
        r1.font.bold = True
        r1.font.color.rgb = RGBColor(27, 54, 93) # Navy
        
        for sub_title, items in subsecs:
            # Heading 2 (Subsection)
            p2 = doc.add_paragraph()
            p2.paragraph_format.left_indent = Inches(0.2)
            p2.paragraph_format.space_before = Pt(4)
            p2.paragraph_format.space_after = Pt(2)
            p2.paragraph_format.keep_with_next = True
            r2 = p2.add_run(sub_title)
            r2.font.name = 'Calibri'
            r2.font.size = Pt(11)
            r2.font.bold = True
            r2.font.color.rgb = RGBColor(0, 90, 156) # Teal
            
            for item in items:
                # Heading 3 (Sub-item)
                p3 = doc.add_paragraph()
                p3.paragraph_format.left_indent = Inches(0.45)
                p3.paragraph_format.space_before = Pt(1)
                p3.paragraph_format.space_after = Pt(1)
                r3 = p3.add_run(f"•  {item}")
                r3.font.name = 'Calibri'
                r3.font.size = Pt(9.5)
                r3.font.color.rgb = RGBColor(50, 50, 50)

if __name__ == "__main__":
    safety_dir = r"c:\Users\Joe\radical_simplicity_ai_os_joe-m\work_in_progress\safety"
    osmm_path = os.path.join(safety_dir, "OSMM_Operations_Safety_Management_Manual_TOC.docx")
    tsmp_path = os.path.join(safety_dir, "TSMP_Transitional_Safety_Management_Plan_TOC.docx")
    
    build_osmm_docx(osmm_path)
    build_tsmp_docx(tsmp_path)
