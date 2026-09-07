import json,os
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
from openpyxl.utils import get_column_letter

base=os.path.dirname(os.path.abspath(__file__))
with open(os.path.join(base,"cri_data_v2.json")) as f: all_data=json.load(f)

cols=[("ID",6),("Plant Area",28),("Critical Task",40),("Role",22),("Phase",16),("Hazards",50),("Consequence",40),("Uncontrolled Risk",14),("Controls & Barriers",55),("Required Competencies",45),("Training & Verification",50),("Residual Risk",14),("Regulatory Reference",40),("ESC Life Saving Rule",22),("SIF Potential",12),("Recommended KPIs",40)]
keys=["Area","Task","Role","Phase","Hazards","Consequence","Pre-Risk","Controls","Competency","Training","Residual","Reg Ref","LSR","SIF","KPI"]

rc={"Critical":PatternFill(start_color="FF0000",end_color="FF0000",fill_type="solid"),"High":PatternFill(start_color="FF8C00",end_color="FF8C00",fill_type="solid"),"Medium":PatternFill(start_color="FFD700",end_color="FFD700",fill_type="solid"),"Low":PatternFill(start_color="32CD32",end_color="32CD32",fill_type="solid")}
rf={"Critical":Font(bold=True,color="FFFFFF",size=10),"High":Font(bold=True,size=10),"Medium":Font(bold=True,size=10),"Low":Font(bold=True,color="FFFFFF",size=10)}
tb=Border(left=Side(style="thin"),right=Side(style="thin"),top=Side(style="thin"),bottom=Side(style="thin"))
hf=PatternFill(start_color="1F3864",end_color="1F3864",fill_type="solid")
hn=Font(bold=True,size=10,color="FFFFFF")
ha=Alignment(horizontal="center",vertical="center",wrap_text=True)

wb=Workbook()
ws=wb.active
ws.title="Critical Risk Inventory"

# Title
ws.merge_cells("A1:P1")
c=ws["A1"];c.value="CRITICAL RISK INVENTORY — LNG (Single Train SMR | Marine Loading | Remote Howe Sound Site)";c.font=Font(bold=True,size=14,color="FFFFFF");c.fill=hf;c.alignment=Alignment(horizontal="center",vertical="center")
ws.row_dimensions[1].height=36

ws.merge_cells("A2:P2")
c=ws["A2"];c.value="Operators & Maintenance Techs | 7/7 or 4/4 Shift | Composite ERT (2FT + Vol) | Passive FF | Helipad + Marine Ambulance | In-House Maintenance | SIMOPS Planned | Squamish Nation SNEAA";c.font=Font(italic=True,size=10,color="FFFFFF");c.fill=PatternFill(start_color="2E5090",end_color="2E5090",fill_type="solid");c.alignment=Alignment(horizontal="center",vertical="center")
ws.row_dimensions[2].height=22

# Headers
for ci,(cn,cw) in enumerate(cols,1):
    c=ws.cell(row=3,column=ci,value=cn);c.font=hn;c.fill=hf;c.alignment=ha;c.border=tb
    ws.column_dimensions[get_column_letter(ci)].width=cw
ws.row_dimensions[3].height=32

# Data
sif_fill=PatternFill(start_color="FF0000",end_color="FF0000",fill_type="solid")
for ri,rec in enumerate(all_data,4):
    ws.cell(row=ri,column=1,value=f"CRI-{ri-3:03d}").font=Font(size=10,bold=True)
    ws.cell(row=ri,column=1).alignment=Alignment(horizontal="center",vertical="top")
    ws.cell(row=ri,column=1).border=tb
    for ci,k in enumerate(keys,2):
        v=rec.get(k,"")
        c=ws.cell(row=ri,column=ci,value=v);c.font=Font(size=10);c.alignment=Alignment(vertical="top",wrap_text=True);c.border=tb
        if k in("Pre-Risk","Residual") and v in rc:
            c.fill=rc[v];c.font=rf[v];c.alignment=Alignment(horizontal="center",vertical="top",wrap_text=True)
        if k=="SIF" and v=="YES":
            c.fill=sif_fill;c.font=Font(bold=True,color="FFFFFF",size=10);c.alignment=Alignment(horizontal="center",vertical="top")
    ws.row_dimensions[ri].height=85

ws.freeze_panes="A4"
ws.auto_filter.ref=f"A3:P{3+len(all_data)}"

# Sheet 2: Risk Matrix
ws2=wb.create_sheet("Risk Matrix Legend")
ws2.merge_cells("A1:E1");ws2["A1"].value="RISK MATRIX";ws2["A1"].font=Font(bold=True,size=14,color="FFFFFF");ws2["A1"].fill=hf
for c2,h in enumerate(["","Catastrophic","Major","Moderate","Minor"],1):
    c=ws2.cell(row=3,column=c2,value=h);c.font=hn;c.fill=hf;c.alignment=Alignment(horizontal="center");c.border=tb;ws2.column_dimensions[get_column_letter(c2)].width=18
matrix=[("Almost Certain",["Critical","Critical","High","High"]),("Likely",["Critical","High","High","Medium"]),("Possible",["High","High","Medium","Medium"]),("Unlikely",["High","Medium","Medium","Low"]),("Rare",["Medium","Medium","Low","Low"])]
for r,(lh,vals) in enumerate(matrix,4):
    c=ws2.cell(row=r,column=1,value=lh);c.font=Font(bold=True,size=10);c.border=tb
    for c2,v in enumerate(vals,2):
        c=ws2.cell(row=r,column=c2,value=v);c.fill=rc.get(v);c.font=rf.get(v);c.alignment=Alignment(horizontal="center");c.border=tb

# ESC Life Saving Rules reference
ws2.cell(row=11,column=1,value="ESC Life Saving Rules:").font=Font(bold=True,size=12)
lsr_list=["LSR-01: Confined Space — Obtain authorization before entry","LSR-02: Energy Isolation — Verify isolation and zero energy before work","LSR-03: Work Authorization — Work with a valid permit when required","LSR-04: Safety Controls — Obtain authorization before overriding/disabling","LSR-05: Working at Height — Protect yourself against a fall","LSR-06: Driving — Follow safe driving rules","LSR-07: Hot Work — Control flammables and ignition sources","LSR-08: Fit for Duty — Be fit for duty (fatigue, impairment)","LSR-09: Safe Mechanical Lifting — Plan and control lifting operations","LSR-10: Line of Fire — Keep yourself out of the line of fire"]
for i,rule in enumerate(lsr_list):
    ws2.cell(row=12+i,column=1,value=rule).font=Font(size=10)

# Sheet 3: KPI Dashboard
ws3=wb.create_sheet("Recommended Safety KPIs")
ws3.merge_cells("A1:E1");ws3["A1"].value="RECOMMENDED SAFETY KPIs — Canadian Oil & Gas Industry Standard";ws3["A1"].font=Font(bold=True,size=14,color="FFFFFF");ws3["A1"].fill=hf
kpi_h=["Category","KPI","Type","Target","Standard/Source"]
kpi_w=[20,45,12,20,30]
for ci,(h,w) in enumerate(zip(kpi_h,kpi_w),1):
    c=ws3.cell(row=3,column=ci,value=h);c.font=hn;c.fill=hf;c.alignment=ha;c.border=tb;ws3.column_dimensions[get_column_letter(ci)].width=w
kpis=[
    ["Process Safety","Tier 1 LOPC Events (API RP 754)","Lagging","Zero","API RP 754; ESC Data Gateway"],
    ["Process Safety","Tier 2 LOPC Events (API RP 754)","Lagging","Zero","API RP 754; ESC Data Gateway"],
    ["Process Safety","SIS/ESD Proof Test Compliance %","Leading","≥98%","IEC 61511"],
    ["Process Safety","Active SIS Bypass Count","Leading","Zero (target)","IEC 61511; ISA-18.2"],
    ["Process Safety","Safety-Critical Maintenance Overdue %","Leading","<2%","API RP 754 Tier 3"],
    ["Process Safety","Alarm Rate (alarms/operator/hour)","Leading","<6 avg, <10 peak","ISA-18.2 (EEMUA 191)"],
    ["Occupational Safety","Total Recordable Injury Frequency (TRIF)","Lagging","<1.0","CAPP; WorkSafeBC"],
    ["Occupational Safety","Lost Time Injury Frequency (LTIF)","Lagging","<0.5","CAPP; WorkSafeBC"],
    ["Occupational Safety","SIF Actual Events","Lagging","Zero","ESC SIF Program"],
    ["Occupational Safety","Potentially Serious Incident (PSI) Reporting Rate","Leading","≥1.0 per 200k hrs","ESC SIF Program"],
    ["Occupational Safety","Life Saving Rule Violations","Lagging","Zero tolerance","ESC Life Saving Rules"],
    ["Occupational Safety","Near-Miss Reporting Rate","Leading","Increasing trend","CAPP"],
    ["Competency","Safety-Critical Training Completion %","Leading","100%","WorkSafeBC; CSA Z276"],
    ["Competency","Emergency Drill Completion vs Schedule %","Leading","≥95%","CSA Z276; NFPA 59A"],
    ["Competency","Competency Assessment Overdue Count","Leading","Zero","CAPP CMS Guidelines"],
    ["Permit to Work","PTW Audit Pass Rate %","Leading","≥95%","OGP Report 394"],
    ["Permit to Work","SIMOPS Conflict Events","Leading","Zero","Site SIMOPS Procedure"],
    ["Marine Operations","ERS Test Compliance %","Leading","100%","SIGTTO; CSA Z276"],
    ["Marine Operations","Marine Loading Incident Count","Lagging","Zero","Transport Canada"],
    ["Environment & Wildlife","Wildlife Incident Count","Lagging","Zero","LNG WMMP; BC Wildlife Act"],
    ["Environment & Wildlife","SNEAA Condition Compliance %","Leading","100%","SNEAA"],
    ["Emergency Response","ERT Drill Completion %","Leading","100%","WorkSafeBC; CSA Z276"],
    ["Emergency Response","Medevac Response Time (min)","Leading","<60 min","Site ERP"],
    ["Contractor/Workforce","D&A Random Test Compliance %","Leading","100%","ESC; API RP 755"],
    ["Contractor/Workforce","Hours-of-Work Exceedance Count","Leading","Zero","API RP 755; WorkSafeBC"],
]
for ri,row in enumerate(kpis,4):
    for ci,v in enumerate(row,1):
        c=ws3.cell(row=ri,column=ci,value=v);c.font=Font(size=10);c.alignment=Alignment(vertical="top",wrap_text=True);c.border=tb
    ws3.row_dimensions[ri].height=28
ws3.freeze_panes="A4"

# Save
out=os.path.join(base,"LNG_LNG_Critical_Risk_Inventory_V2.xlsx")
wb.save(out)
print(f"V2 Excel saved: {out}")
print(f"Total CRI entries: {len(all_data)}")
print(f"Sheets: {wb.sheetnames}")
