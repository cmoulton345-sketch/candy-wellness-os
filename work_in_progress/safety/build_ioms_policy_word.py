import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def add_header_logo_box(doc, doc_id, doc_title):
    header = doc.sections[0].header
    table = header.add_table(rows=1, cols=2, width=Inches(6.5))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    # Left cell: Logo placeholder
    cell_logo = table.cell(0, 0)
    cell_logo.width = Inches(2.2)
    set_cell_background(cell_logo, "F2F4F7")
    set_cell_margins(cell_logo, top=140, bottom=140, left=140, right=140)
    p_logo = cell_logo.paragraphs[0]
    p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_logo = p_logo.add_run("[ INSERT COMPANY LOGO HERE ]")
    run_logo.font.size = Pt(9)
    run_logo.font.bold = True
    run_logo.font.color.rgb = RGBColor(100, 116, 139)
    
    # Right cell: Document metadata
    cell_meta = table.cell(0, 1)
    cell_meta.width = Inches(4.3)
    set_cell_margins(cell_meta, top=100, bottom=100, left=140, right=140)
    p_meta = cell_meta.paragraphs[0]
    p_meta.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run_meta_title = p_meta.add_run(f"{doc_title}\n")
    run_meta_title.font.size = Pt(10)
    run_meta_title.font.bold = True
    run_meta_title.font.color.rgb = RGBColor(15, 23, 42)
    
    run_meta_id = p_meta.add_run(f"Doc ID: {doc_id}  |  Rev: 0  |  CSA Z1000 & Z276 Aligned")
    run_meta_id.font.size = Pt(8.5)
    run_meta_id.font.color.rgb = RGBColor(100, 116, 139)
    
    # Add a styled horizontal line below table in doc body
    p_sep = doc.add_paragraph()
    p_sep.paragraph_format.space_after = Pt(12)
    run_sep = p_sep.add_run("―" * 55)
    run_sep.font.color.rgb = RGBColor(203, 213, 225)
    run_sep.font.bold = True

def style_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    for run in p.runs:
        run.font.name = "Calibri"
        if level == 1:
            run.font.size = Pt(16)
            run.font.color.rgb = RGBColor(30, 58, 138) # Dark Navy
        elif level == 2:
            run.font.size = Pt(13)
            run.font.color.rgb = RGBColor(51, 65, 85)
        elif level == 3:
            run.font.size = Pt(11)
            run.font.color.rgb = RGBColor(71, 85, 105)
    return p

def add_callout_box(doc, text, title="MANAGEMENT REQUIREMENT & MANDATE"):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, "FEF3C7")
    set_cell_margins(cell, top=120, bottom=120, left=160, right=160)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(4)
    r_title = p.add_run(f"⚠ {title}\n")
    r_title.font.bold = True
    r_title.font.size = Pt(10)
    r_title.font.color.rgb = RGBColor(180, 83, 9)
    
    r_text = p.add_run(text)
    r_text.font.size = Pt(9.5)
    r_text.font.color.rgb = RGBColor(120, 53, 15)
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

def build_ioms_policy_doc():
    doc = docx.Document()
    add_header_logo_box(doc, "POL-IOMS-001", "Integrated Operational & Safety Management System (IOMS) Policy")
    
    style_heading(doc, "Integrated Operational & Safety Management System (IOMS) Policy", level=1)
    
    p = doc.add_paragraph()
    p.add_run("Facility Application: ").font.bold = True
    p.add_run("LNG Production, Storage, and Marine Export Facility (Squamish, BC)\n")
    p.add_run("Governing Standards: ").font.bold = True
    p.add_run("CSA Z1000 (OHS Management), CSA Z276 Clause 11 (LNG Operations & Maintenance), WorkSafeBC OHS Regulation, CCPS Risk-Based Process Safety (RBPS)")
    
    style_heading(doc, "1. Executive Policy Statement & Strategic Intent", level=1)
    doc.add_paragraph("LNG establishes that Occupational Health & Safety (SMS) and Process Safety Management (PSM) are structurally embedded into the core Operational Management System (OMS). Safety is not an isolated department or an afterthought; it is the fundamental precondition for every operational hour, turnaround activity, and marine transfer.")
    doc.add_paragraph("This Policy mandates that capital allocation, operating envelopes, maintenance priority, and organizational design must reflect equal weighting between asset protection, human safety, and operational continuity. If operating parameters deviate from established Safe Operating Limits (SOLs) and cannot be stabilized immediately, operations shall be shut down.")
    
    add_callout_box(doc, "EXECUTIVE STOP WORK CHARTER: Every employee and contractor on site possesses absolute authority and mandatory accountability to initiate an immediate Stop Work order or Emergency Shut-in if safety-critical barriers are compromised. Executive Leadership guarantees zero retaliation for safety stops.", "EXECUTIVE SAFETY CHARTER")
    
    style_heading(doc, "2. IOMS Structural Architecture & Integration Model", level=1)
    doc.add_paragraph("The table below illustrates how the tripartite architecture integrates production management (OMS), occupational safety (SMS/OHSMS), and process safety (PSM) across our 7 Operational Pillars:")
    
    table = doc.add_table(rows=8, cols=4)
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    headers = ["The 7 Integrated Pillars", "OMS (Production & Assets)", "SMS / OHSMS (WorkSafeBC / Z1000)", "PSM (CSA Z276 / CCPS)"]
    for i, h in enumerate(headers):
        cell = table.cell(0, i)
        set_cell_background(cell, "1E3A8A")
        set_cell_margins(cell, 100, 100, 120, 120)
        r = cell.paragraphs[0].add_run(h)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        r.font.size = Pt(9)
        
    pillars = [
        ("1. Leadership & Governance", "Capital budget, OPEX allocation, KPI definition, organizational structure", "Safety leadership charter, WCA s.21 compliance, active field walkdowns", "Safety-critical roles matrix, major accident hazard governance"),
        ("2. Risk Assessment", "Operating procedures, production envelope optimization, shift schedules", "Field-Level Hazard Assessment (FLHA), JSA, ergonomics, industrial hygiene", "HAZOP, LOPA, Process Hazard Analysis (PHA), Safe Operating Limits (SOL)"),
        ("3. Safe Work Control", "Daily maintenance planning, outage scheduling, turnaround execution", "Integrated Permit to Work (PTW), Confined Space, WAH, LOTO verification", "Dual-isolation policy (DB&B/Blinding), ignition control, marine safety zones"),
        ("4. Asset Integrity", "Equipment availability, lifecycle management, production throughput", "Workplace hazard guarding, secondary containment, fall arrest gear", "Safety-Critical Elements (SCE) registry, CUI inspection, PM zero-backlog"),
        ("5. Management of Change", "Process optimization, throughput increases, software upgrades", "Worker safety review, ergonomic assessment, PPE adjustment", "Mandatory technical review, Temporary MOC (T-MOC) limits, PSSR sign-off"),
        ("6. Operator Competency", "DCS efficiency, throughput management, equipment handling", "Safety orientation, Life Saving Rules (LSR), fatigue & D&A management", "Control Room Operator (CRO) simulator cert, abnormal scenario drills"),
        ("7. Incident Investigations", "Production loss analysis, equipment downtime tracking", "Recordable injury tracking, WorkSafeBC reporting, near-miss logging", "TapRooT®/ICAM Root Cause Analysis (RCA), HiPo tracking, Corrective Actions")
    ]
    
    for r_idx, row in enumerate(pillars, start=1):
        for c_idx, val in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            set_cell_margins(cell, 80, 80, 100, 100)
            if r_idx % 2 == 0:
                set_cell_background(cell, "F8FAFC")
            p = cell.paragraphs[0]
            r = p.add_run(val)
            r.font.size = Pt(8.5)
            if c_idx == 0:
                r.font.bold = True
                
    doc.add_paragraph().paragraph_format.space_after = Pt(12)
    
    style_heading(doc, "3. Management Level Requirements by Pillar", level=1)
    
    pillar_details = [
        ("Pillar 1: Leadership, Governance & Accountability", [
            "Executive Leadership establishes and maintains the Safety Critical Roles Matrix, explicitly linking WorkSafeBC and CSA Z276 legal accountability to director and superintendent job descriptions.",
            "Mandatory Field Presence: Directors and Superintendents must complete at least 4 documented safety leadership walkdowns per month, focusing on barrier verification and coaching.",
            "Safety performance KPIs (leading and lagging) directly impact executive and management annual compensation incentives."
        ]),
        ("Pillar 2: Risk Identification, Assessment & Control", [
            "No process unit or modification shall be commissioned without an approved Process Hazard Analysis (HAZOP/LOPA) re-validated every 5 years per CSA Z276 Clause 10.",
            "Mandatory SOL Registry: Safe Operating Limits must be hard-coded into the DCS with automated trip setpoints that cannot be overridden by panel operators without formal T-MOC authorization.",
            "Daily Field-Level Hazard Assessments (FLHA) are mandatory prior to commencing any physical task across all shift teams."
        ]),
        ("Pillar 3: Safe Systems of Work (Work Control & Isolation)", [
            "All physical work must be governed by the Electronic Integrated Permit to Work (ePTW) system across Hot Work, Confined Space Entry, and Ground Disturbance.",
            "Zero-Energy Isolation Mandate: All line breaking involving hydrocarbons, cryogenics, or steam requires dual positive isolation (Double Block & Bleed or physical blind/blank).",
            "Marine Exclusion Zones: Automated radar and visual monitoring must enforce a 200-meter safety zone around LNG carriers during berthing and transfer operations."
        ]),
        ("Pillar 4: Asset Integrity & Reliability (SCE Management)", [
            "Safety-Critical Elements (SCEs) such as ESD valves, fire deluge systems, gas detectors, and relief devices must be tracked in the CMMS with designated safety priority tags.",
            "Zero Backlog Policy: Any overdue preventive maintenance on a designated SCE requires immediate notification to the VP Operations and an active mitigation assessment.",
            "Mandatory Corrosion Under Insulation (CUI) regime using pulsed eddy current and profile radiography across all cryogenic and cool-end piping loops."
        ]),
        ("Pillar 5: Management of Change (MOC & PSSR)", [
            "Any physical, procedural, software (DCS interlock), or organizational change requires formal entry into the MOC system prior to execution.",
            "Temporary MOCs (T-MOCs) are strictly time-limited to a maximum of 30 days and must undergo daily shift review by the Shift Charge Engineer.",
            "Pre-Start Safety Review (PSSR): A formal physical walkdown by cross-functional leadership is mandatory before introducing hydrocarbons to any modified asset."
        ]),
        ("Pillar 6: Operational Readiness & Competency", [
            "Control Room Operators (DCS Panel Operators) must achieve certification via the high-fidelity Operator Training Simulator (OTS), completing annual abnormal and emergency shutdown assessments.",
            "Contractor Prequalification: All third-party contractors must maintain an ISNetworld grade of 'A' or 'B' and undergo site-specific cultural and Life Saving Rules orientation.",
            "Fatigue Risk Management System (FRMS): Strict enforcement of shift limits (max 12-hour shifts, minimum 11 hours rest between shifts) and zero-tolerance Drug & Alcohol compliance."
        ]),
        ("Pillar 7: Incident Investigation & Continuous Learning", [
            "All High-Potential (HiPo) near-misses, process safety spills (Tier 1/2 LOPC), and recordable injuries must undergo formal Root Cause Analysis (TapRooT® or ICAM methodology).",
            "Corrective Action Tracking System (CATS): Corrective actions assigned from investigations must be closed within target deadlines; overdue items escalate to executive review weekly.",
            "Lessons Learned dissemination must occur across all shift handovers and safety meetings within 48 hours of any major incident across the global industry."
        ])
    ]
    
    for p_title, bullets in pillar_details:
        style_heading(doc, p_title, level=2)
        for b in bullets:
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_after = Pt(4)
            p.add_run(b).font.size = Pt(9.5)
            
    doc.add_paragraph().paragraph_format.space_after = Pt(16)
    
    # Executive Sign-off Block
    style_heading(doc, "4. Executive Policy Approval & Governance Sign-off", level=1)
    table_sign = doc.add_table(rows=2, cols=2)
    table_sign.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r in range(2):
        for c in range(2):
            cell = table_sign.cell(r, c)
            set_cell_margins(cell, 120, 120, 120, 120)
            set_cell_background(cell, "F8FAFC")
            
    table_sign.cell(0, 0).paragraphs[0].add_run("Approved By: ___________________________\nChief Executive Officer / Managing Director").font.bold = True
    table_sign.cell(0, 1).paragraphs[0].add_run("Endorsed By: ___________________________\nVP Operations & Plant Manager").font.bold = True
    table_sign.cell(1, 0).paragraphs[0].add_run("Effective Date: [ YYYY-MM-DD ]\nDocument ID: POL-IOMS-001")
    table_sign.cell(1, 1).paragraphs[0].add_run("Review Cadence: Annual / Post-Turnaround\nNext Review Date: [ YYYY-MM-DD ]")
    
    doc.save("POL-IOMS-001_Integrated_OMS_SMS_Policy.docx")
    print("IOMS Policy Word document generated successfully.")

if __name__ == "__main__":
    build_ioms_policy_doc()
